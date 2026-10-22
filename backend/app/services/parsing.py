"""Document parsing: PDF (PyMuPDF), DOCX (python-docx), TXT and Markdown.

Every parser returns pages whose text is the exact string later sliced into chunks, so chunk
offsets, page numbers and quotes can always be traced back to the source page.
"""
from __future__ import annotations

import io
import re
import statistics
import zipfile
from dataclasses import dataclass, field
from datetime import date

ALLOWED_TYPES = {"pdf", "docx", "txt", "md"}
EXTENSION_MAP = {".pdf": "pdf", ".docx": "docx", ".txt": "txt", ".md": "md", ".markdown": "md"}


class ParseError(Exception):
    """Unrecoverable problem with the file itself (empty, corrupt, encrypted, unsupported)."""


class NeedsOCR(Exception):
    """The document contains no extractable text layer (e.g. a scanned PDF)."""


@dataclass
class ParsedPage:
    page_number: int
    text: str
    line_offset: int = 0  # source line number (0-based) of the first line in this page
    headings: list[tuple[int, str]] = field(default_factory=list)  # (char offset, heading text)


@dataclass
class ParsedDocument:
    pages: list[ParsedPage]
    page_mode: str  # physical (PDF) | explicit_breaks | rendered_breaks | form_feed | none
    title: str | None = None
    author: str | None = None
    published_date: date | None = None
    metadata: dict = field(default_factory=dict)
    warnings: list[str] = field(default_factory=list)

    @property
    def total_chars(self) -> int:
        return sum(len(p.text.strip()) for p in self.pages)


def detect_file_type(filename: str, head: bytes) -> str:
    ext = ""
    m = re.search(r"(\.[A-Za-z0-9]+)$", filename or "")
    if m:
        ext = m.group(1).lower()
    ftype = EXTENSION_MAP.get(ext)
    if ftype is None:
        raise ParseError(f"不支持的文件类型 '{ext or '无扩展名'}'，仅支持 PDF、DOCX、TXT、Markdown")
    if ftype == "pdf" and not head.startswith(b"%PDF"):
        raise ParseError("文件扩展名为 .pdf，但内容不是有效的 PDF（缺少 %PDF 文件头）")
    if ftype == "docx" and not head.startswith(b"PK"):
        raise ParseError("文件扩展名为 .docx，但内容不是有效的 DOCX（不是 ZIP 容器）")
    return ftype


# ---------------------------------------------------------------- dates
_DATE_PATTERNS = [
    re.compile(r"(?:发布日期|发布时间|日期|修订日期|Date|Published)\s*[:：]?\s*(\d{4})[-/.年](\d{1,2})[-/.月](\d{1,2})"),
]


def find_date_in_text(text: str) -> date | None:
    for pat in _DATE_PATTERNS:
        m = pat.search(text[:3000])
        if m:
            try:
                return date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
            except ValueError:
                continue
    return None


def _parse_iso_date(value: str | None) -> date | None:
    if not value:
        return None
    m = re.match(r"\D*(\d{4})[-/.]?(\d{2})[-/.]?(\d{2})", str(value))
    if not m:
        return None
    try:
        return date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    except ValueError:
        return None


def _assign_date(doc: ParsedDocument, meta_date: date | None):
    text_date = find_date_in_text(doc.pages[0].text if doc.pages else "")
    if text_date:
        doc.published_date = text_date
        doc.metadata["published_date_source"] = "document_text"
    elif meta_date:
        doc.published_date = meta_date
        doc.metadata["published_date_source"] = "file_metadata"


# ---------------------------------------------------------------- PDF
def parse_pdf(data: bytes) -> ParsedDocument:
    import pymupdf

    try:
        pdf = pymupdf.open(stream=data, filetype="pdf")
    except Exception as exc:  # noqa: BLE001 - PyMuPDF raises several types
        raise ParseError(f"PDF 文件损坏或无法打开：{exc}") from exc
    if pdf.needs_pass:
        raise ParseError("PDF 已加密，需要密码才能读取")
    if pdf.page_count == 0:
        raise ParseError("PDF 不包含任何页面")

    pages: list[ParsedPage] = []
    image_only_pages: list[int] = []
    try:
        for index, page in enumerate(pdf):
            text, headings = _pdf_page_text(page)
            if len(text.strip()) < 5 and page.get_images():
                image_only_pages.append(index + 1)
            pages.append(ParsedPage(page_number=index + 1, text=text, headings=headings))
    except Exception as exc:  # noqa: BLE001
        raise ParseError(f"PDF 内容解析失败（文件可能损坏）：{exc}") from exc

    meta = pdf.metadata or {}
    doc = ParsedDocument(
        pages=pages,
        page_mode="physical",
        title=(meta.get("title") or None),
        author=(meta.get("author") or None),
        metadata={k: v for k, v in meta.items() if v},
    )
    # Use the TOC (bookmarks) as section headings where available.
    try:
        for level, title, page_no in pdf.get_toc():
            if 1 <= page_no <= len(pages):
                p = pages[page_no - 1]
                idx = p.text.find(title.strip())
                if idx >= 0 and all(h[1] != title.strip() for h in p.headings):
                    p.headings.append((idx, title.strip()))
                    p.headings.sort()
    except Exception:  # noqa: BLE001 - TOC is optional
        pass

    if doc.total_chars == 0:
        if image_only_pages:
            raise NeedsOCR(
                f"PDF 共 {len(pages)} 页，均为图像且没有可提取的文本层（可能是扫描件），需要先进行 OCR 处理"
            )
        raise ParseError("PDF 中没有可提取的文本")
    if image_only_pages:
        doc.warnings.append(f"第 {', '.join(map(str, image_only_pages))} 页只有图像、无文本层，需 OCR 才能检索")
        doc.metadata["pages_needing_ocr"] = image_only_pages
    _assign_date(doc, _parse_iso_date(meta.get("creationDate")))
    return doc


def _pdf_page_text(page) -> tuple[str, list[tuple[int, str]]]:
    """Rebuild page text from the layout dict so we know where headings start."""
    layout = page.get_text("dict", sort=True)
    lines: list[tuple[str, float, bool]] = []  # (text, font size, starts new block)
    for block in layout.get("blocks", []):
        if block.get("type") != 0:
            continue
        first = True
        for line in block.get("lines", []):
            spans = line.get("spans", [])
            text = "".join(s.get("text", "") for s in spans).rstrip()
            if not text.strip():
                continue
            size = max((s.get("size", 0) for s in spans), default=0)
            lines.append((text, size, first))
            first = False
    if not lines:
        return "", []
    sizes = [s for _, s, _ in lines]
    body = statistics.median(sizes)
    parts: list[str] = []
    headings: list[tuple[int, str]] = []
    offset = 0
    for i, (text, size, new_block) in enumerate(lines):
        sep = "" if i == 0 else ("\n\n" if new_block else "\n")
        offset += len(sep)
        parts.append(sep + text)
        if size >= body * 1.2 and len(text.strip()) <= 80:
            headings.append((offset, text.strip()))
        offset += len(text)
    return "".join(parts), headings


# ---------------------------------------------------------------- DOCX
_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def parse_docx(data: bytes) -> ParsedDocument:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            if "word/document.xml" not in zf.namelist():
                raise ParseError("DOCX 文件缺少 word/document.xml，文件可能损坏")
        document = docx.Document(io.BytesIO(data))
    except ParseError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise ParseError(f"DOCX 文件损坏或无法打开：{exc}") from exc

    has_explicit = any(
        br.get(f"{_W}type") == "page" for br in document.element.body.iter(f"{_W}br")
    ) or any(True for _ in document.element.body.iter(f"{_W}pageBreakBefore"))
    has_rendered = any(True for _ in document.element.body.iter(f"{_W}lastRenderedPageBreak"))
    mode = "explicit_breaks" if has_explicit else ("rendered_breaks" if has_rendered else "none")

    pages: list[list[str]] = [[]]
    headings: list[list[tuple[int, str]]] = [[]]

    def page_len(i: int) -> int:
        return len("\n\n".join(pages[i]))

    def new_page():
        pages.append([])
        headings.append([])

    def breaks_in(element) -> int:
        n = 0
        for el in element.iter():
            if mode == "explicit_breaks":
                if el.tag == f"{_W}br" and el.get(f"{_W}type") == "page":
                    n += 1
            elif mode == "rendered_breaks" and el.tag == f"{_W}lastRenderedPageBreak":
                n += 1
        return n

    for child in document.element.body.iterchildren():
        if child.tag == f"{_W}p":
            para = Paragraph(child, document)
            if mode == "explicit_breaks" and child.find(f".//{_W}pageBreakBefore") is not None and pages[-1]:
                new_page()
            text = para.text.strip()
            style = (para.style.name if para.style is not None else "") or ""
            if text:
                is_heading = style.lower().startswith(("heading", "title")) or style.startswith(("标题",))
                if is_heading:
                    pos = page_len(len(pages) - 1) + (2 if pages[-1] else 0)
                    headings[-1].append((pos, text))
                pages[-1].append(text)
            for _ in range(breaks_in(child)):
                new_page()
        elif child.tag == f"{_W}tbl":
            table = Table(child, document)
            rows = []
            for row in table.rows:
                cells = [c.text.strip().replace("\n", " ") for c in row.cells]
                rows.append(" | ".join(cells))
            if rows:
                pages[-1].append("\n".join(rows))

    parsed_pages = []
    for i, paras in enumerate(pages):
        text = "\n\n".join(paras)
        if not text.strip() and i == len(pages) - 1 and parsed_pages:
            continue  # trailing break
        parsed_pages.append(ParsedPage(page_number=len(parsed_pages) + 1, text=text, headings=headings[i]))

    props = document.core_properties
    doc = ParsedDocument(
        pages=parsed_pages,
        page_mode=mode,
        title=props.title or None,
        author=props.author or None,
        metadata={"docx_page_mode": mode},
    )
    if mode == "none":
        doc.warnings.append("DOCX 未包含分页信息，页码统一记为第 1 页（逻辑页）")
    if doc.total_chars == 0:
        raise ParseError("DOCX 中没有可提取的文本")
    _assign_date(doc, props.created.date() if props.created else None)
    return doc


# ---------------------------------------------------------------- TXT / MD
def decode_text(data: bytes) -> str:
    if b"\x00" in data[:4096]:
        raise ParseError("文本文件包含二进制内容（NUL 字节），可能不是文本文件或已损坏")
    for enc in ("utf-8-sig", "gb18030"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    raise ParseError("无法识别文本编码（支持 UTF-8 / GB18030）")


def _split_form_feed(text: str) -> tuple[list[ParsedPage], str]:
    raw_pages = text.split("\f")
    pages = []
    line = 0
    for i, chunk in enumerate(raw_pages):
        pages.append(ParsedPage(page_number=i + 1, text=chunk, line_offset=line))
        line += chunk.count("\n")
    return pages, ("form_feed" if len(raw_pages) > 1 else "none")


def parse_txt(data: bytes) -> ParsedDocument:
    text = decode_text(data).replace("\r\n", "\n").replace("\r", "\n")
    if not text.strip():
        raise ParseError("文本文件内容为空")
    pages, mode = _split_form_feed(text)
    first_line = next((ln.strip() for ln in text.splitlines() if ln.strip()), None)
    doc = ParsedDocument(pages=pages, page_mode=mode, title=(first_line[:200] if first_line else None))
    if mode == "none":
        doc.warnings.append("纯文本没有分页信息，页码统一记为第 1 页，引用以行号定位")
    _assign_date(doc, None)
    return doc


_FRONTMATTER = re.compile(r"\A---\n(.*?)\n---\n", re.S)
_MD_HEADING = re.compile(r"^(#{1,6})\s+(.+?)\s*#*\s*$", re.M)


def parse_markdown(data: bytes) -> ParsedDocument:
    text = decode_text(data).replace("\r\n", "\n").replace("\r", "\n")
    if not text.strip():
        raise ParseError("Markdown 文件内容为空")
    meta: dict = {}
    m = _FRONTMATTER.match(text)
    if m:
        for line in m.group(1).splitlines():
            if ":" in line:
                k, v = line.split(":", 1)
                meta[k.strip().lower()] = v.strip().strip("'\"")
        # Blank out frontmatter with the same number of lines so line numbers stay exact.
        text = "\n" * m.group(0).count("\n") + text[m.end():]
    pages, mode = _split_form_feed(text)
    for p in pages:
        p.headings = [(hm.start(), hm.group(2).strip()) for hm in _MD_HEADING.finditer(p.text)]
    title = meta.get("title")
    if not title:
        h1 = next((h for p in pages for h in p.headings), None)
        title = h1[1] if h1 else None
    doc = ParsedDocument(
        pages=pages, page_mode=mode, title=title, author=meta.get("author"), metadata={"frontmatter": meta}
    )
    if mode == "none":
        doc.warnings.append("Markdown 没有分页信息，页码统一记为第 1 页，引用以行号定位")
    if not doc.total_chars:
        raise ParseError("Markdown 文件除元数据外没有正文")
    fm_date = _parse_iso_date(meta.get("date") or meta.get("published"))
    if fm_date:
        doc.published_date = fm_date
        doc.metadata["published_date_source"] = "frontmatter"
    else:
        _assign_date(doc, None)
    return doc


def parse_document(file_type: str, data: bytes) -> ParsedDocument:
    if not data:
        raise ParseError("文件为空（0 字节）")
    if file_type == "pdf":
        return parse_pdf(data)
    if file_type == "docx":
        return parse_docx(data)
    if file_type == "txt":
        return parse_txt(data)
    if file_type == "md":
        return parse_markdown(data)
    raise ParseError(f"不支持的文件类型：{file_type}")
