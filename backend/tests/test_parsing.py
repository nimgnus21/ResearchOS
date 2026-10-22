"""1. PDF/DOCX/TXT/MD parsing, 3. empty & corrupted files."""
import io
from datetime import date

import pymupdf
import pytest
from docx import Document
from docx.enum.text import WD_BREAK

from app.services.parsing import NeedsOCR, ParseError, detect_file_type, parse_document


def test_pdf_pages_headings_and_metadata(corpus_dir):
    doc = parse_document("pdf", (corpus_dir / "kx200_revised_report_2024.pdf").read_bytes())
    assert doc.page_mode == "physical"
    assert [p.page_number for p in doc.pages] == [1, 2, 3, 4]
    assert "0.35 mm/100 mm" in doc.pages[1].text
    assert "GB/T 18254" in doc.pages[2].text
    assert "400 运行小时" in doc.pages[3].text
    assert doc.pages[1].headings[0][1] == "二、根本原因"
    assert doc.title == "KX-200 轴承过热事故修订调查报告"
    assert doc.published_date == date(2024, 6, 20)
    assert doc.metadata["published_date_source"] == "document_text"


def test_docx_page_breaks_headings_and_tables(corpus_dir):
    doc = parse_document("docx", (corpus_dir / "kx200_vendor_bulletin.docx").read_bytes())
    assert doc.page_mode == "explicit_breaks"
    assert len(doc.pages) == 2
    assert "40°C" in doc.pages[0].text and "PAO" not in doc.pages[0].text
    assert "合成润滑脂" in doc.pages[1].text
    assert "高温（>40°C） | PAO 合成脂 | 300 小时" in doc.pages[1].text  # table preserved
    assert any(t == "润滑建议" for _, t in doc.pages[1].headings)
    assert doc.published_date == date(2024, 2, 10)


def test_docx_without_breaks_is_single_logical_page():
    d = Document()
    d.add_paragraph("第一段。")
    d.add_paragraph("第二段。")
    buf = io.BytesIO()
    d.save(buf)
    doc = parse_document("docx", buf.getvalue())
    assert doc.page_mode == "none" and len(doc.pages) == 1
    assert any("页码统一记为第 1 页" in w for w in doc.warnings)


def test_txt_form_feed_pages_and_line_offsets(corpus_dir):
    doc = parse_document("txt", (corpus_dir / "operation_log_summary.txt").read_bytes())
    assert doc.page_mode == "form_feed"
    assert len(doc.pages) == 2
    assert "7.8 mm/s" in doc.pages[0].text
    assert "62°C" in doc.pages[1].text
    assert doc.pages[1].line_offset == doc.pages[0].text.count("\n")


def test_txt_gb18030_and_plain():
    doc = parse_document("txt", "轴承温度正常。".encode("gb18030"))
    assert doc.pages[0].text == "轴承温度正常。"
    assert doc.page_mode == "none"


def test_markdown_frontmatter_preserves_line_numbers(corpus_dir):
    raw = (corpus_dir / "maintenance_team_memo.md").read_text(encoding="utf-8")
    doc = parse_document("md", raw.encode())
    assert doc.title == "维护班组关于 KX-200 轴承过热的意见备忘"
    assert doc.published_date == date(2024, 7, 5)
    assert doc.author == "三号泵站维护班组（虚构）"
    # Frontmatter is blanked but line count preserved -> line numbers match the original file.
    assert doc.pages[0].text.count("\n") == raw.count("\n")
    assert "---" not in doc.pages[0].text
    assert [t for _, t in doc.pages[0].headings][:2] == ["维护班组意见备忘", "结论"]


@pytest.mark.parametrize("ftype", ["pdf", "docx", "txt", "md"])
def test_empty_file_rejected(ftype):
    with pytest.raises(ParseError, match="为空"):
        parse_document(ftype, b"")


def test_whitespace_only_text_rejected():
    with pytest.raises(ParseError, match="为空"):
        parse_document("txt", b"   \n\n  ")


def test_corrupted_pdf():
    with pytest.raises(ParseError, match="损坏|无法"):
        parse_document("pdf", b"%PDF-1.7\n this is not really a pdf \x00\x01\x02")


def test_corrupted_docx():
    with pytest.raises(ParseError, match="损坏|缺少"):
        parse_document("docx", b"PK\x03\x04garbage-not-a-zip")


def test_binary_masquerading_as_text():
    with pytest.raises(ParseError, match="二进制"):
        parse_document("txt", b"abc\x00\x00\x00def")


def test_extension_content_mismatch():
    with pytest.raises(ParseError, match="不是有效的 PDF"):
        detect_file_type("report.pdf", b"hello wo")
    with pytest.raises(ParseError, match="不支持"):
        detect_file_type("image.png", b"\x89PNG")


def test_scanned_pdf_requires_ocr(corpus_dir):
    with pytest.raises(NeedsOCR, match="OCR"):
        parse_document("pdf", (corpus_dir / "scanned_inspection_sheet.pdf").read_bytes())


def test_encrypted_pdf():
    d = pymupdf.open()
    d.new_page().insert_text((50, 50), "secret")
    data = d.tobytes(encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw="pw", owner_pw="pw")
    with pytest.raises(ParseError, match="加密"):
        parse_document("pdf", data)


def test_pdf_with_mixed_scanned_page_warns():
    d = pymupdf.open()
    d.new_page().insert_text((50, 50), "Page one has real text about bearings.")
    p2 = d.new_page()
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 50, 50), False)
    pix.clear_with(200)
    p2.insert_image(pymupdf.Rect(0, 0, 100, 100), pixmap=pix)
    doc = parse_document("pdf", d.tobytes())
    assert doc.metadata["pages_needing_ocr"] == [2]
    assert any("OCR" in w for w in doc.warnings)


def test_docx_page_break_helper_roundtrip():
    d = Document()
    d.add_paragraph("A 页内容").runs[0].add_break(WD_BREAK.PAGE)
    d.add_paragraph("B 页内容").runs[0].add_break(WD_BREAK.PAGE)
    d.add_paragraph("C 页内容")
    buf = io.BytesIO()
    d.save(buf)
    doc = parse_document("docx", buf.getvalue())
    assert [p.text for p in doc.pages] == ["A 页内容", "B 页内容", "C 页内容"]
